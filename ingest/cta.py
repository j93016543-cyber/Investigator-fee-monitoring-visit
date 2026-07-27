"""
CTA(기관별 임상시험계약) 폴더 파서 - 하위그룹 A 소스
=====================================================

data/source/cta/Site <번호>/<버전 폴더>/... 구조를 스캔해서 site 별 CTA 버전
(version, effective_date, protocol_version, institution, PI, 항목별 비용)을 추출한다.

- version        : 버전 폴더명 (Initial / Amd1_PA V7.0 / CTA AMD#2_PA8 ...)
- effective_date : 폴더 내 파일명에서 파싱한 가장 늦은 날짜(= 서명/FE 일자)
- items          : budget 엑셀의 (항목, Selected Cost) 목록

새 site 는 data/source/cta/ 에 'Site <번호>' 폴더째 넣으면 자동 인식된다.
"""
import re
import datetime
from pathlib import Path

import openpyxl

MONTHS = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
DATE_RE = re.compile(r"(\d{1,2})(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)(\d{4}|\d{2})", re.I)


def parse_dates(text):
    out = []
    for d, mon, y in DATE_RE.findall(text):
        yr = int(y);  yr += 2000 if yr < 100 else 0
        try:
            out.append(datetime.date(yr, MONTHS[mon.lower()], int(d)))
        except ValueError:
            pass
    return out


def _meta(ws):
    """budget 시트 상단 메타(Protocol Version / Institution / PI)."""
    m = {}
    for r in range(1, 14):
        for c in range(1, 4):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            key = v.strip().rstrip(":")
            nxt = ws.cell(r, c + 1).value
            if key == "Protocol Version" and nxt:
                m["protocol_version"] = str(nxt).strip()
            elif key == "Institution" and nxt:
                m["institution"] = str(nxt).strip()
            elif key == "Principal Investigator" and nxt:
                m["pi"] = str(nxt).strip()
    return m


def extract_items(path):
    """budget 엑셀에서 (항목, Selected Cost) 추출. 두 레이아웃 지원."""
    items = {}
    meta = {}
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except Exception:
        return items, meta
    for ws in wb.worksheets:
        if not meta:
            meta = _meta(ws)
        # 레이아웃 1: 헤더에 'Selected Cost' + 'Trial Procedures'
        hdr_row = cost_col = name_col = None
        for r in range(1, min(ws.max_row, 20) + 1):
            for c in range(1, min(ws.max_column, 34) + 1):
                v = ws.cell(r, c).value
                if isinstance(v, str):
                    if "Selected Cost" in v:
                        hdr_row, cost_col = r, c
                    elif "Trial Procedures" in v:
                        name_col = c
            if hdr_row and name_col:
                break
        if hdr_row and cost_col and name_col:
            for r in range(hdr_row + 1, ws.max_row + 1):
                nm = ws.cell(r, name_col).value
                cost = ws.cell(r, cost_col).value
                if isinstance(nm, str) and nm.strip() and isinstance(cost, (int, float)):
                    items.setdefault(nm.strip(), round(float(cost), 2))
        # 레이아웃 2: 'Overhead Percent' 있는 요약시트 → col A 항목 + col B 숫자
        elif any("Overhead Percent" in str(ws.cell(r, 1).value or "")
                 for r in range(1, min(ws.max_row, 12) + 1)):
            for r in range(13, ws.max_row + 1):
                a = ws.cell(r, 1).value
                b = ws.cell(r, 2).value
                if isinstance(a, str) and a.strip() and isinstance(b, (int, float)):
                    items.setdefault(a.strip(), round(float(b), 2))
    wb.close()
    return items, meta


def parse_site_folder(site_dir):
    """'Site <번호>' 폴더 → ctas 리스트."""
    ctas = []
    for ver_dir in sorted([d for d in site_dir.iterdir() if d.is_dir()]):
        files = list(ver_dir.rglob("*"))
        # effective date = 폴더 내 파일명 최신 날짜
        dates = []
        for f in files:
            dates += parse_dates(f.name)
        eff = max(dates).isoformat() if dates else None
        # 항목: budget 엑셀에서 추출 (여러 개면 병합)
        items = {}
        meta = {}
        for f in files:
            if f.suffix.lower() == ".xlsx":
                it, mt = extract_items(f)
                items.update(it)
                if mt and not meta:
                    meta = mt
        ctas.append({
            "version": ver_dir.name,
            "effective_date": eff,
            "protocol_version": meta.get("protocol_version"),
            "institution": meta.get("institution"),
            "pi": meta.get("pi"),
            "items": [{"item": k, "amount": v} for k, v in items.items()],
            "n_files": len([f for f in files if f.is_file()]),
        })
    # effective_date 순 정렬
    ctas.sort(key=lambda c: c["effective_date"] or "")
    return ctas


def parse_cta_root(cta_root):
    """data/source/cta/ 전체 스캔 → {site#: {name, ctas}}."""
    out = {}
    if not cta_root.exists():
        return out
    for site_dir in sorted(cta_root.iterdir()):
        if not site_dir.is_dir():
            continue
        m = re.search(r"(\d{3,4})", site_dir.name)
        if not m:
            continue
        site = m.group(1)
        ctas = parse_site_folder(site_dir)
        name = None
        for c in ctas:
            if c.get("institution"):
                name = c["institution"]
                break
        out[site] = {"name": name, "ctas": ctas, "source_folder": site_dir.name}
    return out
