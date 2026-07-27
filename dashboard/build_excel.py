#!/usr/bin/env python3
"""
data/store.json + data/site_cta.json  ->  exports/TTK-CS-101_tracker.xlsx
대시보드(A~F)를 엑셀 워크북으로 출력 (음영 + 통화별 표기).
사용법: python3 dashboard/build_excel.py
"""
import json
import re
import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
store = json.loads((ROOT / "data" / "store.json").read_text(encoding="utf-8"))
sct_path = ROOT / "data" / "site_cta.json"
scts = json.loads(sct_path.read_text(encoding="utf-8")).get("sites", {}) if sct_path.exists() else {}
C = store["contract"]

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="0D7680")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
SEC_FILL = PatternFill("solid", fgColor="D9E4E6")
SUB_FILL = PatternFill("solid", fgColor="EDF2F3")
TOT_FILL = PatternFill("solid", fgColor="CFDBDE")
EFF_FILL = PatternFill("solid", fgColor="E2F0F1")
ROWH_FILL = PatternFill("solid", fgColor="EDF2F3")
SHADE = {"good": PatternFill("solid", fgColor="D8EEE4"),
         "warn": PatternFill("solid", fgColor="FBECCB"),
         "crit": PatternFill("solid", fgColor="F7DDD7"),
         "info": PatternFill("solid", fgColor="E4EBF6")}
THIN = Side(style="thin", color="D0D8DA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BASE = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)


def real_date(x):
    return x if (x and not str(x).startswith("2000-01-01")) else ""


def parse_visit(v):
    if not v:
        return ("", 99999)
    if re.search(r"screen", v, re.I):
        return ("Screening", 0)
    m = re.search(r"C(\d+).*?D\s*([0-9]+(?:_[0-9A-Za-z]+)?)", v, re.I)
    if m:
        cyc = int(m.group(1)); dn = int(re.match(r"\d+", m.group(2)).group())
        return (f"C{cyc}D{m.group(2)}", cyc * 1000 + dn)
    return (v[:16], 98000)


def cta_list(site):
    return (scts.get(str(site), {}) or {}).get("ctas", [])


def cta_at(site, date):
    arr = sorted([c for c in cta_list(site) if c.get("effective_date")], key=lambda c: c["effective_date"])
    if not arr:
        return None
    if not date:
        return arr[-1]
    out = None
    for c in arr:
        if c["effective_date"] <= date:
            out = c
    return out or arr[0]


def site_name(s):
    return (store.get("sites", {}).get(str(s)) or {}).get("name", "")


def site_cur(s):
    o = store.get("sites", {}).get(str(s)) or {}
    return "KRW" if (o.get("country") == "South Korea") else "USD"


def money_fmt(cur):
    return '₩#,##0' if cur == "KRW" else '$#,##0.00'


def visit_cost_match(actual, costs):
    if not costs or not actual:
        return None
    aa = abs(actual); mn = costs[0]
    if any(abs(aa - x) <= max(1, x * 0.02) for x in costs):
        return "good"
    if aa < mn * 0.98:
        return "crit"
    return "warn"


def numfmt(cell, kind):
    cell.number_format = {"usd": '#,##0.00', "usd0": '#,##0', "int": '#,##0', "krw": '₩#,##0',
                          "pct": '0.0"%"'}.get(kind, 'General')


def sheet(wsname, headers, rows, widths=None, numcols=None, money_cols=None, freeze="A2"):
    """money_cols: {colidx: cur_header}  → 해당 셀을 행의 통화로 ₩/$ 표기."""
    ws = wb.create_sheet(wsname)
    ws.sheet_view.showGridLines = False
    numcols = numcols or {}
    money_cols = money_cols or {}
    for j, h in enumerate(headers, 1):
        c = ws.cell(1, j, h); c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(rows, 2):
        rcls = row.get("_cls"); shade = row.get("_shade", {})
        for j, h in enumerate(headers, 1):
            cell = ws.cell(i, j, row.get(h))
            cell.font = BOLD if rcls in ("sec", "total") else BASE
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="right" if (j in numcols or j in money_cols) else "left",
                                       vertical="center")
            if j in numcols:
                numfmt(cell, numcols[j])
            if j in money_cols:
                cur = row.get(money_cols[j]) or "USD"
                cell.number_format = money_fmt(cur)
            fill = {"sec": SEC_FILL, "sub": SUB_FILL, "total": TOT_FILL, "eff": EFF_FILL}.get(rcls)
            if fill:
                cell.fill = fill
            if shade.get(h):
                cell.fill = SHADE[shade[h]]
    ws.freeze_panes = freeze
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return ws


wb = openpyxl.Workbook()
fees = store.get("investigator_fees", [])
invs = store.get("pass_through_invoices", [])
pays = store.get("payments", [])
env = C.get("payment_envelopes", {})
sites_cta = sorted([s for s in scts if scts[s].get("ctas")], key=str)

# ============================================================ 0. 요약
ws = wb.active; ws.title = "0. 요약"; ws.sheet_view.showGridLines = False
feeU = sum(x["amount"] for x in fees if x.get("currency") == "USD" and isinstance(x.get("amount"), (int, float)))
feeK = sum(x["amount"] for x in fees if x.get("currency") == "KRW" and isinstance(x.get("amount"), (int, float)))
ws["A1"] = "TTK-CS-101 · 연구비 & 모니터링 누적 트래커"
ws["A1"].font = Font(name=FONT, size=16, bold=True, color="0D7680")
ws["A2"] = f"Protocol {C.get('protocol','TTK-CS-101')} · CRO {C.get('vendor','')} · 계약번호 {C.get('contract_number','')}"
ws["A2"].font = Font(name=FONT, size=10, color="6D8388")
ws["A3"] = f"유효 계약: {C.get('effective_version')} (eff. {C.get('effective_date')})  ·  한국 site 비용은 ₩(WON) 표기"
ws["A3"].font = Font(name=FONT, size=10, italic=True)
kpis = [
    ("연구비(Investigator Grants) 계약 총액", env.get("investigator_grants_total"), "usd0", "WO v3 (USD)"),
    ("연구비 실지급 누계 (USD, 미국 site)", round(feeU, 2), "usd", "visit_activity! Visit Amount"),
    ("연구비 실지급 누계 (₩ WON, 한국 site)", round(feeK), "krw", "visit_activity! Visit Amount"),
    ("Invoiceable(Professional) 최대", env.get("professional_specialty_max"), "usd0", "WO v3"),
    ("현재 IQVIA 계약 예산(최신 CO)", (C.get("cnf_versions") or [{}])[-1].get("grand_total"), "usd0", "IQVIA CNF 최신"),
    ("연구비 실지급 건수", len(fees), "int", "visit_activity rows"),
]
ws.cell(5, 1, "항목").font = HDR_FONT; ws.cell(5, 1).fill = HDR_FILL
ws.cell(5, 2, "값").font = HDR_FONT; ws.cell(5, 2).fill = HDR_FILL
ws.cell(5, 3, "출처").font = HDR_FONT; ws.cell(5, 3).fill = HDR_FILL
for k, (lab, val, kind, src) in enumerate(kpis):
    rr = 6 + k
    ws.cell(rr, 1, lab).font = BASE
    c = ws.cell(rr, 2, val); c.font = BOLD; numfmt(c, kind)
    ws.cell(rr, 3, src).font = Font(name=FONT, size=9, color="6D8388")
    for cc in range(1, 4):
        ws.cell(rr, cc).border = BORDER

# 분기별 지급 (연구비 vs invoice 분리) — req2
def q_of(dt):
    if not dt:
        return None
    y, m = int(dt[:4]), int(dt[5:7])
    return f"{y} Q{(m-1)//3+1}"

qsplit = {}
for x in pays:
    qk = q_of(x.get("invoiced_date") or x.get("paid_date"))
    amt = x.get("invoiced_usd")
    if not qk or not isinstance(amt, (int, float)):
        continue
    d2 = qsplit.setdefault(qk, {"direct": 0.0, "ptc": 0.0})
    d2["direct" if x.get("cost_type") == "Direct" else "ptc"] += amt
r0 = 6 + len(kpis) + 2
ws.cell(r0, 1, "분기별 IQVIA 지급 — 연구비(Direct) vs invoice(PTC/Pass-through), USD").font = Font(name=FONT, size=12, bold=True)
r0 += 1
for j, h in enumerate(["연도-분기", "연구비(Direct)", "invoice(PTC)", "합계"], 1):
    c = ws.cell(r0, j, h); c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER
for k, qk in enumerate(sorted(qsplit)):
    rr = r0 + 1 + k; o = qsplit[qk]
    ws.cell(rr, 1, qk).font = BASE
    for jj, v in [(2, o["direct"]), (3, o["ptc"]), (4, o["direct"] + o["ptc"])]:
        c = ws.cell(rr, jj, round(v, 2)); c.font = BASE; c.number_format = '#,##0'
    for cc in range(1, 5):
        ws.cell(rr, cc).border = BORDER

# 연구비 지급일별 요약 — req5
r0 = r0 + len(qsplit) + 3
ws.cell(r0, 1, "연구비 지급일별 요약 (visit_activity Payment Date)").font = Font(name=FONT, size=12, bold=True)
r0 += 1
for j, h in enumerate(["지급일(Paid date)", "건수", "USD 합", "₩(WON) 합"], 1):
    c = ws.cell(r0, j, h); c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER
pdmap = {}
for x in fees:
    pd = x.get("payment_date")
    if not pd:
        continue
    o = pdmap.setdefault(pd, {"n": 0, "u": 0.0, "k": 0.0})
    o["n"] += 1
    if x.get("currency") == "KRW":
        o["k"] += x.get("amount") or 0
    else:
        o["u"] += x.get("amount") or 0
for k, pd in enumerate(sorted(pdmap)):
    rr = r0 + 1 + k; o = pdmap[pd]
    ws.cell(rr, 1, pd).font = BASE
    ws.cell(rr, 2, o["n"]).font = BASE
    cu = ws.cell(rr, 3, round(o["u"], 2)); cu.number_format = '$#,##0.00'
    ck = ws.cell(rr, 4, round(o["k"])); ck.number_format = '₩#,##0'
    for cc in range(1, 5):
        ws.cell(rr, cc).border = BORDER
ws.column_dimensions["A"].width = 44
for col in "BCD":
    ws.column_dimensions[col].width = 20
ws.freeze_panes = "A5"

# ============================================================ A. CTA 변경이력
rowsA = []
for s in sites_cta:
    rowsA.append({"CTA Version": f"Site {s} — {site_name(s) or scts[s].get('name','')}", "_cls": "sec"})
    cs = sorted(scts[s]["ctas"], key=lambda c: c.get("effective_date") or "")
    eff = cta_at(s, None); prev = None
    for c in cs:
        n = len(c.get("items", []))
        delta = "" if prev is None else (f"+{n-prev}" if n - prev >= 0 else str(n - prev))
        rowsA.append({"CTA Version": c["version"], "Effective Date": c.get("effective_date") or "—",
                      "Protocol Ver.": c.get("protocol_version") or "—", "항목수": n or "—", "Δ항목": delta,
                      "_cls": "eff" if (eff and c["version"] == eff["version"]) else None})
        prev = n
sheet("A. CTA 변경이력", ["CTA Version", "Effective Date", "Protocol Ver.", "항목수", "Δ항목"],
      rowsA, widths=[44, 16, 28, 10, 10], numcols={4: "int"})

# ============================================================ A2. CTA 항목별비용
rowsAI = []
for s in sites_cta:
    cur = site_cur(s)
    for c in sorted(scts[s]["ctas"], key=lambda c: c.get("effective_date") or ""):
        for it in c.get("items", []):
            rowsAI.append({"Site#": s, "Cur": cur, "CTA Version": c["version"],
                           "Effective": c.get("effective_date") or "—",
                           "항목(Trial Procedure)": it.get("item"), "Selected Cost": it.get("amount")})
sheet("A2. CTA 항목별비용", ["Site#", "Cur", "CTA Version", "Effective", "항목(Trial Procedure)", "Selected Cost"],
      rowsAI, widths=[7, 6, 26, 12, 60, 16], money_cols={6: "Cur"})

# ============================================================ B. 지급원장 (계약대조 포함 — req1,3)
rowsB = []
for x in fees:
    lab, _ = parse_visit(x.get("visit")); vd = real_date(x.get("visit_date"))
    c = cta_at(x.get("site"), vd or x.get("payment_date"))
    amt = x.get("amount"); neg = isinstance(amt, (int, float)) and amt < 0
    costs = c.get("visit_costs", {}).get(lab) if c else None
    if not c:
        chk, shade = "계약없음", "warn"
    elif costs is None:
        chk, shade = "계약외 절차", "warn"   # 계약/프로토콜에 없는 방문·절차 (req3)
    else:
        m = visit_cost_match(amt, costs)
        chk = {"good": "일치", "warn": "초과", "crit": "미달"}.get(m, "—")
        shade = m
    rowsB.append({"Country": x.get("country"), "site name": site_name(x.get("site")), "site #": x.get("site"),
                  "구분": "Investigator fee", "subject #": x.get("patient"), "Visit #": lab, "Visit date": vd,
                  "description": x.get("visit"), "Amount": amt, "Cur": x.get("currency"),
                  "CTA 계약금액": (f"{costs[0]:,.0f}~{costs[-1]:,.0f}" if costs and costs[0] != costs[-1]
                                 else (f"{costs[0]:,.0f}" if costs else "—")),
                  "계약대조": chk, "effective CTA": c["version"] if c else "—",
                  "CTA eff date": c["effective_date"] if c else "—",
                  "Payment #": x.get("payment_no"), "Paid date": x.get("payment_date"),
                  "비고": "조정/취소(음수)" if neg else "",
                  "_shade": {"Amount": "warn" if neg else None, "계약대조": shade}})
for x in invs:
    c = cta_at(x.get("site"), x.get("payment_date"))
    rowsB.append({"Country": x.get("country"), "site name": site_name(x.get("site")), "site #": x.get("site"),
                  "구분": "Invoice", "subject #": "", "Visit #": "", "Visit date": "",
                  "description": x.get("description"), "Amount": x.get("amount"), "Cur": x.get("currency"),
                  "CTA 계약금액": "—", "계약대조": "invoice(대조제외)", "effective CTA": c["version"] if c else "—",
                  "CTA eff date": c["effective_date"] if c else "—",
                  "Payment #": x.get("payment_no"), "Paid date": x.get("payment_date"), "비고": "",
                  "_shade": {}})
Bhead = ["Country", "site name", "site #", "구분", "subject #", "Visit #", "Visit date", "description",
         "Amount", "Cur", "CTA 계약금액", "계약대조", "effective CTA", "CTA eff date", "Payment #", "Paid date", "비고"]
sheet("B. 지급원장", Bhead, rowsB,
      widths=[11, 32, 7, 15, 12, 9, 11, 30, 14, 6, 16, 14, 22, 12, 11, 11, 16],
      money_cols={9: "Cur"})

# ============================================================ C. 대상자·방문 매트릭스 (template 형식 — req4)
idx = {}
for x in fees:
    s, pt = x.get("site"), x.get("patient")
    if not s or not pt:
        continue
    lab, order = parse_visit(x.get("visit"))
    o = idx.setdefault((s, pt, lab), {"site": s, "pt": pt, "label": lab, "cur": x.get("currency"),
                                      "u": 0.0, "k": 0.0, "dates": [], "vdate": ""})
    o["order"] = order
    if x.get("currency") == "KRW":
        o["k"] += x.get("amount") or 0
    else:
        o["u"] += x.get("amount") or 0
    if x.get("payment_date"):
        o["dates"].append(x["payment_date"])
    if not o["vdate"]:
        o["vdate"] = real_date(x.get("visit_date"))
# per site->patient->[visit cells]
tree = {}
for o in idx.values():
    tree.setdefault(str(o["site"]), {}).setdefault(o["pt"], []).append(o)
wsC = wb.create_sheet("C. 대상자·방문"); wsC.sheet_view.showGridLines = False
wsC.column_dimensions["A"].width = 22
for col in "BCDEFGHIJKLMNOPQRSTU":
    wsC.column_dimensions[col].width = 14
METRICS = ["Visit date", "effective CTA", "CTA effective date",
           "Investigator fee amount", "Invoiceable item amount", "Paid date"]
rr = 1
for s in sorted(tree, key=str):
    scur = tree[s]  # patients
    for pt in sorted(scur):
        cells = sorted(scur[pt], key=lambda o: o["order"])
        cur = cells[0]["cur"] or site_cur(s)
        # 헤더(대상자)
        wsC.cell(rr, 1, f"Site {s} — {site_name(s)}").font = Font(name=FONT, size=10, bold=True, color="0D7680")
        rr += 1
        hc = wsC.cell(rr, 1, "Visit #"); hc.font = HDR_FONT; hc.fill = HDR_FILL; hc.border = BORDER
        wsC.cell(rr, 2, pt).font = BOLD  # 대상자 표기
        for j, o in enumerate(cells, 3):
            cc = wsC.cell(rr, j, o["label"]); cc.font = HDR_FONT; cc.fill = HDR_FILL; cc.border = BORDER
            cc.alignment = Alignment(horizontal="center")
        rr += 1
        for mi, metric in enumerate(METRICS):
            lc = wsC.cell(rr, 1, metric); lc.font = BOLD; lc.fill = ROWH_FILL; lc.border = BORDER
            for j, o in enumerate(cells, 3):
                vd = o["vdate"] or (sorted(o["dates"])[0] if o["dates"] else "")
                c = cta_at(s, vd)
                actual = o["k"] if cur == "KRW" else o["u"]
                costs = c.get("visit_costs", {}).get(o["label"]) if c else None
                m = visit_cost_match(actual, costs)
                cell = wsC.cell(rr, j); cell.border = BORDER; cell.font = BASE
                cell.alignment = Alignment(horizontal="right")
                if metric == "Visit date":
                    cell.value = vd or "—"
                elif metric == "effective CTA":
                    cell.value = c["version"] if c else "—"
                elif metric == "CTA effective date":
                    cell.value = c["effective_date"] if c else "—"
                elif metric == "Investigator fee amount":
                    if actual:
                        cell.value = round(actual, 2)
                        cell.number_format = money_fmt(cur)
                        if m:
                            cell.fill = SHADE[m]
                    else:
                        cell.value = "—"
                elif metric == "Invoiceable item amount":
                    cell.value = "—"  # site 단위(방문 매핑 없음)
                elif metric == "Paid date":
                    ds = sorted([d for d in o["dates"] if d])
                    cell.value = ds[-1] if ds else "—"
            rr += 1
        rr += 1  # blank
wsC.freeze_panes = "B1"

# ============================================================ CRO 지급원장 (IQVIA Direct/PTC — req2)
rowsP = []
rowsP.append({"구분": "■ Direct (연구비/Professional) — 지출결의# 포함", "_cls": "sec"})
for x in [p for p in pays if p.get("cost_type") == "Direct"]:
    rowsP.append({"구분": "Direct", "Milestone/설명": x.get("milestone"), "Invoice #": x.get("invoice_no"),
                  "Invoiced date": x.get("invoiced_date"), "Paid date": x.get("paid_date"),
                  "Amount(USD)": x.get("invoiced_usd"), "총액(₩)": x.get("total_krw"),
                  "지출결의#(내부기안)": x.get("approval_no"), "지급": "Y" if x.get("paid_yn") == "Y" else ""})
rowsP.append({"구분": "■ PTC (Pass-through/invoice) — invoice#·date", "_cls": "sec"})
for x in [p for p in pays if p.get("cost_type") == "PTC"]:
    rowsP.append({"구분": "PTC", "Milestone/설명": "Pass-through invoice", "Invoice #": x.get("invoice_no"),
                  "Invoiced date": x.get("invoiced_date"), "Paid date": "",
                  "Amount(USD)": x.get("invoiced_usd"), "총액(₩)": "", "지출결의#(내부기안)": "", "지급": ""})
sheet("CRO지급(IQVIA)",
      ["구분", "Milestone/설명", "Invoice #", "Invoiced date", "Paid date", "Amount(USD)", "총액(₩)", "지출결의#(내부기안)", "지급"],
      rowsP, widths=[10, 40, 14, 14, 14, 16, 16, 18, 6], numcols={6: "usd", 7: "int"})

# ============================================================ D. IQVIA 계약
rowsD = []
rowsD.append({"항목": "■ 계약 버전 이력", "_cls": "sec"})
for v in C.get("versions", []):
    rowsD.append({"항목": v["version"], "Unit/유형": v.get("type"), "Effective": v.get("effective_date"),
                  "금액/총액(USD)": v.get("grand_total"), "비고": v.get("note"),
                  "_cls": "eff" if v.get("is_effective") else None})
rowsD.append({"항목": "■ IQVIA CO/CNF 변경이력 (Revised Grand Total)", "_cls": "sec"})
prev = None
for v in C.get("cnf_versions", []):
    g = v.get("grand_total")
    delta = "" if (prev is None or g is None) else f"Δ {g-prev:+,.0f}"
    rowsD.append({"항목": v["version"], "Effective": v.get("effective_date"), "금액/총액(USD)": g, "비고": delta})
    if g is not None:
        prev = g
sheet("D. IQVIA 계약", ["항목", "Unit/유형", "Effective", "금액/총액(USD)", "비고"], rowsD,
      widths=[46, 14, 14, 18, 40], numcols={4: "usd"})

# ============================================================ D2. 지급스케줄
rowsPS = [{"예정월": x["month"], "Milestone": x["milestone"], "%": x["pct"], "Net Amount(USD)": x["amount"],
           "_cls": "total" if x["milestone"] == "Total" else None} for x in C.get("payment_schedule", [])]
sheet("D2. IQVIA 지급스케줄", ["예정월", "Milestone", "%", "Net Amount(USD)"], rowsPS,
      widths=[10, 52, 8, 18], numcols={3: "pct", 4: "usd0"})

# ============================================================ D3. 계약 항목·비용 (CNF, 단가변경 — req6)
cnf = {v["version"]: v for v in C.get("cnf_versions", [])}
c2 = {x["item"]: x for x in cnf.get("CNF 2", {}).get("budget", [])}
c3 = cnf.get("CNF 3", {}).get("budget", [])
MONI = ("Monitoring", "Initiation", "Qualification", "Close", "Interim", "Co-Monitoring", "Travel", "Onsite")
rowsD3 = []
for x in c3:
    kind = x.get("kind")
    u3 = x.get("unit_cost"); u2 = (c2.get(x["item"]) or {}).get("unit_cost")
    chg = None
    if isinstance(u3, (int, float)) and isinstance(u2, (int, float)) and abs(u3 - u2) > 0.005:
        chg = f"{u2:,.2f} → {u3:,.2f}"
    is_moni = any(k in x["item"] for k in MONI)
    rowsD3.append({"항목(Project Cost)": x["item"], "Unit": x.get("unit"), "수량": x.get("qty"),
                   "단가(USD, CNF3)": u3, "Revised 총액": x.get("total"),
                   "CNF2→CNF3 변경": chg or ("" if kind == "item" else ""),
                   "_cls": ("sec" if kind == "section" else "subtotal" if kind == "subtotal" else "total" if kind == "total" else None),
                   "_shade": {"단가(USD, CNF3)": "warn"} if chg else ({"항목(Project Cost)": "info"} if is_moni and kind == "item" else {})})
sheet("D3. 계약항목·비용(CNF)",
      ["항목(Project Cost)", "Unit", "수량", "단가(USD, CNF3)", "Revised 총액", "CNF2→CNF3 변경"], rowsD3,
      widths=[54, 20, 8, 16, 16, 22], numcols={3: "int", 4: "usd", 5: "usd0"})

# ============================================================ E. CRA 경비
cvis = store.get("cra_visits", [])
cidx = {}
for v in cvis:
    cidx.setdefault(re.sub(r"[\s,]", "", str(v.get("cra", "")).lower()), []).append(v.get("visit_start"))
uc = C.get("monitoring_unit_costs", {})


def fmatch(x):
    arr = cidx.get(re.sub(r"[\s,]", "", str(x.get("cra", "")).lower()))
    if not arr:
        return "warn"
    dt = x.get("visit_date") or x.get("trans_date")
    for vs in arr:
        try:
            if dt and vs and abs((datetime.date.fromisoformat(vs) - datetime.date.fromisoformat(dt)).days) <= 10:
                return "good"
        except Exception:
            pass
    return "warn"


rowsE = []
for x in store.get("cra_expenses", []):
    amt = x.get("net_amount") if x.get("net_amount") is not None else x.get("amount")
    m = fmatch(x)
    trav = x.get("group") == "교통비" and isinstance(amt, (int, float)) and amt > uc.get("Monitoring_Travel_per_visit", 1e9)
    rowsE.append({"CRA": x.get("cra"), "Site#": x.get("site"), "카테고리": x.get("group"), "Reason": x.get("reason"),
                  "Description": x.get("description"), "Trans Date": x.get("trans_date"), "Visit Date": x.get("visit_date"),
                  "금액(USD)": amt, "F일치": "일치" if m == "good" else "미확인",
                  "_shade": {"금액(USD)": "warn" if trav else None, "F일치": m}})
sheet("E. CRA경비", ["CRA", "Site#", "카테고리", "Reason", "Description", "Trans Date", "Visit Date", "금액(USD)", "F일치"],
      rowsE, widths=[16, 7, 12, 10, 34, 12, 12, 14, 10], numcols={8: "usd"})

# ============================================================ F. CRA visit
rowsF = []
for v in sorted(cvis, key=lambda v: (str(v.get("cra")), str(v.get("site")), str(v.get("visit_start")))):
    rowsF.append({"CRA": v.get("cra"), "Site#": v.get("site"), "기관(Account)": v.get("account"), "PI": v.get("pi"),
                  "Visit Type": v.get("visit_type"), "Status": v.get("status"), "Visit Start": v.get("visit_start"),
                  "Visit End": v.get("visit_end"), "DOS": v.get("dos"), "Report": v.get("report_status")})
sheet("F. CRA visit",
      ["CRA", "Site#", "기관(Account)", "PI", "Visit Type", "Status", "Visit Start", "Visit End", "DOS", "Report"],
      rowsF, widths=[16, 7, 34, 20, 22, 16, 12, 12, 6, 12], numcols={9: "int"})

# ============================================================ 모니터링 잔여 (+F 집계 참고 — req7)
mv = store.get("monitoring_visits", {})
snaps = mv.get("snapshots", [])
rowsM = []
for r in mv.get("rows", []):
    row = {"카테고리": ("" if r.get("activity") else r.get("group")) or "",
           "Activity": r.get("activity") or r.get("group") or "", "계약": r.get("contracted"), "잔여": r.get("balance")}
    for i, a in enumerate(r.get("actuals", [])):
        if i < len(snaps):
            row[snaps[i]] = a.get("total")
    bal = r.get("balance")
    row["_shade"] = {"잔여": ("crit" if (isinstance(bal, (int, float)) and bal < 0) else
                             ("warn" if bal == 0 else ("good" if isinstance(bal, (int, float)) else None)))}
    if str(r.get("group") or "").startswith(("Interim MV Sub-Total", "Total")):
        row["_cls"] = "total"
    rowsM.append(row)
# F(CRA report) visit type 집계 — 참고
from collections import Counter
fcount = Counter(v.get("visit_type") for v in cvis if v.get("visit_type"))
rowsM.append({"카테고리": "", "Activity": "", "계약": ""})
rowsM.append({"카테고리": "■ 참고: F(CRA Site Visit Report) 실적 집계 (별도 소스)", "_cls": "sec"})
for vt, n in fcount.most_common():
    rowsM.append({"카테고리": "F실적", "Activity": vt, "잔여": n})
Mhead = ["카테고리", "Activity", "계약"] + snaps + ["잔여"]
numM = {3: "int"}
for i in range(len(snaps)):
    numM[4 + i] = "int"
numM[4 + len(snaps)] = "int"
sheet("모니터링 잔여", Mhead, rowsM, widths=[30, 40, 8] + [11] * len(snaps) + [8], numcols=numM)

(ROOT / "exports").mkdir(exist_ok=True)
out = ROOT / "exports" / "TTK-CS-101_tracker.xlsx"
wb.save(out)
print("saved:", out)
print("sheets:", wb.sheetnames)
